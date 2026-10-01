#!/usr/bin/env python3
"""Apply a reviewed catalog operation to the canonical workbook, then validate it."""
import argparse
import json
import re
import tempfile
from datetime import date
from pathlib import Path
from copy import copy
from importlib.machinery import SourceFileLoader
from openpyxl import load_workbook
from catalog_core import validate_catalog
importer = SourceFileLoader("activity_importer", str(Path(__file__).resolve().parent / "import-activity.py")).load_module()
validate_input, append_with_style, slugify, ImportErrorMessage = importer.validate_input, importer.append_with_style, importer.slugify, importer.ImportErrorMessage

ROOT = Path(__file__).resolve().parents[1]
generator = SourceFileLoader('catalog_generator', str(ROOT / 'scripts/generate-catalog.py')).load_module()
BOOK = ROOT / 'data/catalogo-actividades.xlsx'

def cells(sheet):
    keys = {c.value: c.column for c in sheet[1]}
    return [(n, {k: sheet.cell(n, col).value for k, col in keys.items()}) for n in range(2, sheet.max_row + 1) if sheet.cell(n, 1).value is not None]

def put(sheet, row, fields):
    keys = {c.value: c.column for c in sheet[1]}
    for k, value in fields.items():
        if k not in keys: raise ImportErrorMessage(f'Falta la columna {k} en {sheet.title}')
        sheet.cell(row, keys[k], value)

def remove(sheet, rows):
    for row in sorted(rows, reverse=True): sheet.delete_rows(row)

def one(sheet, key, value):
    hits = [(n, row) for n, row in cells(sheet) if row[key] == value]
    if len(hits) != 1: raise ImportErrorMessage(f'No existe el registro {value}')
    return hits[0]

def clean(value, limit=180, optional=False):
    if not isinstance(value, str) or len(value) > limit or any(ord(c) < 32 for c in value) or (not optional and not value.strip()):
        raise ImportErrorMessage('Texto incompleto o inválido')
    return value.strip()

def image_fields(w, payload, kind):
    image = payload.get('image')
    if not image: return {}
    if not isinstance(image, dict) or set(image) != {'path', 'alt', 'author', 'source', 'license'}:
        raise ImportErrorMessage('Metadatos de imagen incompletos')
    path = image['path']
    if not isinstance(path, str) or not re.fullmatch(r'/images/admin/(topics|activities)/[a-f0-9]{32}\.(png|jpg|webp)', path) or kind not in path:
        raise ImportErrorMessage('Ruta de imagen inválida')
    if not (ROOT / 'public' / path.lstrip('/')).is_file(): raise ImportErrorMessage('La imagen aún no está almacenada en el repositorio')
    alt = clean(image['alt']); author = clean(image['author']); source = clean(image['source'], 300); license_name = clean(image['license'])
    credit_id = 'credit-admin-' + Path(path).stem
    if not any(row['ID'] == credit_id for _, row in cells(w['CRÉDITOS'])):
        append_with_style(w['CRÉDITOS'], {'ID':credit_id, 'Título':alt, 'Autoría':author, 'Fuente':source, 'Licencia':license_name})
    return {'Imagen':path, 'Texto alternativo':alt, 'ID de crédito':credit_id}

def apply(w, data, catalog, today):
    action = data.get('action'); entity = data.get('entity'); item = data.get('item')
    if action not in {'create','update','delete'} or entity not in {'activity','topic'} or not isinstance(item, dict):
        raise ImportErrorMessage('Operación inválida')
    if entity == 'activity':
        sheet = w['ACTIVIDADES']; links = w['ACTIVIDAD_TEMA']
        if action == 'create':
            record = validate_input(item, catalog)
            append_with_style(sheet, {'ID':record['id'], 'Slug':record['slug'], 'Título':record['title'], 'Título de origen':record['sourceTitle'], 'Descripción':record['description'], 'ID de asignatura principal':record['subject'], 'IDs de asignatura':record['subject'], 'ID de tema principal':record['topics'][0], 'Idioma':record['language'], 'ID de tipo':record['type'], 'Fecha de publicación':today, 'Fecha de actualización':today, 'Estado editorial':'published', 'Clase de origen':'external', 'ID de plataforma':record['platform'], 'ID de recurso':record['resource'], 'URL pública':record['url'], 'URL canónica':record['url'], 'URL original':record['url'], 'Estado del enlace':'verified' if record['verified'] else 'unverified', 'Última verificación':today if record['verified'] else None, 'Título verificado':bool(record['verified'] and record['sourceTitle']), 'Fuente del título':'Página pública de la actividad' if record['sourceTitle'] else 'Entrada manual', 'Fuente de la descripción':'Entrada manual' if record['description'] else None, **image_fields(w, item, 'activities')})
            for topic in record['topics']: append_with_style(links, {'ID de actividad':record['id'],'ID de tema':topic})
            return record['id']
        ident = clean(item.get('id'))
        n, old = one(sheet, 'ID', ident)
        if old['Clase de origen'] != 'external': raise ImportErrorMessage('Esta actividad nativa es de solo lectura desde el panel; edita el Excel y su contenido en el repositorio')
        if action == 'delete':
            if ident in {a['id'] for a in catalog['activities'] if a.get('legacy',{}).get('subjectLabel') == 'Ciencias / Conocimiento del Medio'}: raise ImportErrorMessage('La actividad pertenece a la migración histórica de Ciencias')
            if any(ident in str(row.get('IDs de actividad relacionada') or '').split('|') for _,row in cells(w['RECURSOS_RELACIONADOS'])):
                raise ImportErrorMessage('Un recurso multimedia referencia esta actividad. Edítalo antes de borrar.')
            if any(row['ID de actividad'] == ident for _,row in cells(w['CONTROL_DE_CALIDAD'])):
                raise ImportErrorMessage('Hay incidencias de calidad asociadas. Revísalas antes de borrar.')
            remove(links,[i for i,row in cells(links) if row['ID de actividad']==ident]);sheet.delete_rows(n)
            return ident
        title = clean(item.get('title')); desc = clean(item.get('description',''),1200,True)
        subject = clean(item.get('subjectId')); topics = item.get('topicIds')
        if subject not in {x['id'] for x in catalog['subjects'] if x['status']=='active'} or not isinstance(topics,list) or not topics or len(topics)!=len(set(topics)) or len(topics)>20 or any(t not in {x['id'] for x in catalog['topics'] if x['subjectId']==subject and x['status']=='active'} for t in topics): raise ImportErrorMessage('Asignatura o temas inválidos')
        typ = clean(item.get('typeId'));language=clean(item.get('language'),35)
        if typ not in {x['id'] for x in catalog['activityTypes']} or not re.fullmatch(r'[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*', language): raise ImportErrorMessage('Tipo o idioma inválido')
        # IDs, slug, URL and historical dates are intentionally immutable in an edit.
        put(sheet,n,{'Título':title,'Descripción':desc,'ID de asignatura principal':subject,'IDs de asignatura':subject,'ID de tema principal':topics[0],'Idioma':language,'ID de tipo':typ,'Fecha de actualización':today,'Fuente de la descripción':'Entrada manual' if desc else None, **image_fields(w,item,'activities')})
        remove(links,[i for i,row in cells(links) if row['ID de actividad']==ident])
        for topic in topics: append_with_style(links,{'ID de actividad':ident,'ID de tema':topic})
        return ident
    sheet = w['TEMAS']; ident = item.get('id')
    if action == 'create':
        subject = clean(item.get('subjectId')); name=clean(item.get('name')); description=clean(item.get('description',''),1200,True)
        if subject not in {x['id'] for x in catalog['subjects'] if x['status']=='active'}: raise ImportErrorMessage('Asignatura inválida')
        slug=slugify(name,'').rstrip('-');ident=f'topic-{subject}-{slug}'
        if any(t['id']==ident or t['slug']==slug or slug in t.get('legacySlugs',[]) for t in catalog['topics']): raise ImportErrorMessage('Ya existe un tema con ese nombre o slug')
        order=max([t['order'] for t in catalog['topics'] if t['subjectId']==subject] or [0])+1
        append_with_style(sheet,{'ID':ident,'ID de asignatura':subject,'Nombre':name,'Slug':slug,'Orden':order,'Descripción':description,'Estado':'active',**image_fields(w,item,'topics')})
        return ident
    ident=clean(ident);n,old=one(sheet,'ID',ident)
    if action == 'update':
        fields={'Nombre':clean(item.get('name')),'Descripción':clean(item.get('description',''),1200,True),**image_fields(w,item,'topics')}
        put(sheet,n,fields);return ident
    if any(row['Nuevo ID de tema']==ident for _,row in cells(w['MIGRACIÓN_TEMAS'])): raise ImportErrorMessage('Este tema conserva una ruta histórica y no se puede borrar')
    refs=[(i,row) for i,row in cells(w['ACTIVIDAD_TEMA']) if row['ID de tema']==ident]
    media=[row for _,row in cells(w['RECURSOS_RELACIONADOS']) if ident in str(row.get('IDs de tema') or '').split('|') or row.get('ID de tema principal')==ident]
    if media: raise ImportErrorMessage('Este tema tiene recursos multimedia asociados. Reasígnalos primero en el Excel.')
    replacement=item.get('replacementTopicId')
    if refs:
        valid={t['id'] for t in catalog['topics'] if t['subjectId']==old['ID de asignatura'] and t['id']!=ident and t['status']=='active'}
        if replacement not in valid: raise ImportErrorMessage('Reasigna las actividades a otro tema activo de la misma asignatura')
        for _,row in refs:
            activity_id=row['ID de actividad']; _,a=one(w['ACTIVIDADES'],'ID',activity_id)
            if any(x['ID de actividad']==activity_id and x['ID de tema']==replacement for _,x in cells(w['ACTIVIDAD_TEMA'])): raise ImportErrorMessage('La reasignación duplicaría una relación de actividad')
        for i,_ in refs: put(w['ACTIVIDAD_TEMA'],i,{'ID de tema':replacement})
        for i,a in cells(w['ACTIVIDADES']):
            if a['ID de tema principal']==ident: put(w['ACTIVIDADES'],i,{'ID de tema principal':replacement,'Fecha de actualización':today})
    sheet.delete_rows(n)
    return ident

def main():
    p=argparse.ArgumentParser();p.add_argument('--payload',type=Path,required=True);p.add_argument('--input',type=Path,default=BOOK);p.add_argument('--output',type=Path,default=BOOK);p.add_argument('--date',default=date.today().isoformat());a=p.parse_args()
    try:
        data=json.loads(a.payload.read_text());catalog,_=generator.load_general_workbook(a.input);w=load_workbook(a.input)
        ident=apply(w,data,catalog,a.date)
        with tempfile.NamedTemporaryFile(suffix='.xlsx',dir=a.output.parent,delete=False) as f: tmp=Path(f.name)
        try:
            w.save(tmp);generated,locations=generator.load_general_workbook(tmp);errors,_=validate_catalog(generated,locations)
            if errors: raise ImportErrorMessage('Validación fallida: '+'; '.join(e['problem'] for e in errors[:5]))
            tmp.replace(a.output)
        finally:tmp.unlink(missing_ok=True)
        print(json.dumps({'id':ident,'action':data['action'],'entity':data['entity']}))
    except (ValueError,TypeError,KeyError) as e: raise SystemExit(str(e)) from e
if __name__=='__main__': main()
